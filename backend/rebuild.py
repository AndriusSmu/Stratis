import os
import shutil

# Define all files and their content
files = {
    "app/__init__.py": '''"""Stratis Backend - Marketing Campaign Manager"""

__version__ = "2.0.0"
''',

    "app/config.py": '''
"""
Configuration management for Stratis backend.
"""

import os
from dataclasses import dataclass, field
from typing import List, Optional
from pathlib import Path


@dataclass
class DatabaseConfig:
    """Database configuration."""
    path: str = field(default_factory=lambda: os.getenv("STRATIS_DB_PATH", "campaigns.db"))
    
    @property
    def absolute_path(self) -> Path:
        return Path(self.path).resolve()


@dataclass
class AIConfig:
    """AI/Ollama configuration."""
    url: str = field(default_factory=lambda: os.getenv("STRATIS_OLLAMA_URL", "http://localhost:11434/api/generate"))
    model: str = field(default_factory=lambda: os.getenv("STRATIS_OLLAMA_MODEL", "llama3.2"))
    timeout: int = field(default_factory=lambda: int(os.getenv("STRATIS_OLLAMA_TIMEOUT", "60")))
    enabled: bool = field(default_factory=lambda: os.getenv("STRATIS_AI_ENABLED", "true").lower() == "true")


@dataclass
class APIConfig:
    """API configuration."""
    prefix: str = field(default_factory=lambda: os.getenv("STRATIS_API_PREFIX", ""))
    environment: str = field(default_factory=lambda: os.getenv("STRATIS_ENVIRONMENT", "development"))
    debug: bool = field(default_factory=lambda: os.getenv("STRATIS_DEBUG", "false").lower() == "true")
    host: str = field(default_factory=lambda: os.getenv("STRATIS_HOST", "127.0.0.1"))
    port: int = field(default_factory=lambda: int(os.getenv("STRATIS_PORT", "8000")))
    reload: bool = field(default_factory=lambda: os.getenv("STRATIS_RELOAD", "false").lower() == "true")


@dataclass
class CORSConfig:
    """CORS configuration."""
    origins: List[str] = field(default_factory=lambda: os.getenv("STRATIS_CORS_ORIGINS", "*").split(","))
    allow_credentials: bool = True
    allow_methods: List[str] = field(default_factory=lambda: ["*"])
    allow_headers: List[str] = field(default_factory=lambda: ["*"])


@dataclass
class RateLimitConfig:
    """Rate limiting configuration."""
    requests: int = field(default_factory=lambda: int(os.getenv("STRATIS_RATE_LIMIT_REQUESTS", "100")))
    period: int = field(default_factory=lambda: int(os.getenv("STRATIS_RATE_LIMIT_PERIOD", "60")))


@dataclass
class AuthConfig:
    """Authentication configuration."""
    secret_key: str = field(default_factory=lambda: os.getenv("STRATIS_SECRET_KEY", "dev-secret-key-change-me"))
    algorithm: str = field(default_factory=lambda: os.getenv("STRATIS_JWT_ALGORITHM", "HS256"))
    access_token_expire_minutes: int = field(default_factory=lambda: int(os.getenv("STRATIS_ACCESS_TOKEN_EXPIRE_MINUTES", "30")))


@dataclass
class Config:
    """Main configuration container."""
    database: DatabaseConfig = field(default_factory=DatabaseConfig)
    ai: AIConfig = field(default_factory=AIConfig)
    api: APIConfig = field(default_factory=APIConfig)
    cors: CORSConfig = field(default_factory=CORSConfig)
    rate_limit: RateLimitConfig = field(default_factory=RateLimitConfig)
    auth: AuthConfig = field(default_factory=AuthConfig)
    
    @classmethod
    def from_env(cls) -> "Config":
        return cls()
    
    def ensure_directories(self) -> None:
        db_dir = self.database.absolute_path.parent
        if not db_dir.exists():
            db_dir.mkdir(parents=True, exist_ok=True)


config = Config.from_env()
''',

    "app/database/__init__.py": '''"""Database module for Stratis backend."""

from .connection import DatabaseConnection, get_db
from .repositories import CampaignRepository, UserRepository

__all__ = ["DatabaseConnection", "get_db", "CampaignRepository", "UserRepository"]
''',

    "app/database/connection.py": '''
"""
Database connection management.
"""

import sqlite3
import threading
from contextlib import contextmanager
from typing import Iterator
from pathlib import Path

from ..config import config


class DatabaseConnection:
    """Thread-safe database connection manager."""
    
    _instance = None
    _lock = threading.Lock()
    
    def __new__(cls):
        if cls._instance is None:
            with cls._lock:
                if cls._instance is None:
                    cls._instance = super().__new__(cls)
                    cls._instance._initialized = False
        return cls._instance
    
    def __init__(self):
        if self._initialized:
            return
        self._initialized = True
        self._db_path = config.database.absolute_path
        self._local = threading.local()
        self._initialize_db()
    
    def _initialize_db(self) -> None:
        conn = self._get_connection()
        try:
            cursor = conn.cursor()
            
            # Users table
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    email TEXT UNIQUE NOT NULL,
                    username TEXT UNIQUE NOT NULL,
                    password_hash TEXT NOT NULL,
                    full_name TEXT,
                    is_active INTEGER DEFAULT 1,
                    created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT DEFAULT CURRENT_TIMESTAMP
                )
            """)
            
            # Campaigns table with user_id
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS campaigns (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL,
                    name TEXT NOT NULL,
                    description TEXT,
                    budget REAL NOT NULL,
                    spent REAL DEFAULT 0,
                    currency TEXT DEFAULT 'USD',
                    start_date TEXT NOT NULL,
                    end_date TEXT,
                    target_audience TEXT,
                    status TEXT DEFAULT 'Draft',
                    owner TEXT,
                    tags TEXT,
                    assets TEXT,
                    notes TEXT,
                    created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY (user_id) REFERENCES users (id) ON DELETE CASCADE
                )
            """)
            
            # Indexes
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_campaigns_user_id ON campaigns(user_id)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_campaigns_status ON campaigns(status)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_campaigns_start_date ON campaigns(start_date)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_campaigns_end_date ON campaigns(end_date)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_users_username ON users(username)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_users_email ON users(email)")
            
            conn.commit()
        finally:
            conn.close()
    
    def _get_connection(self) -> sqlite3.Connection:
        if not hasattr(self._local, "connection"):
            self._local.connection = sqlite3.connect(
                str(self._db_path),
                check_same_thread=False,
                timeout=30.0
            )
            self._local.connection.row_factory = sqlite3.Row
        return self._local.connection
    
    @contextmanager
    def get_cursor(self) -> Iterator[sqlite3.Cursor]:
        conn = self._get_connection()
        cursor = conn.cursor()
        try:
            yield cursor
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            cursor.close()
    
    def execute(self, query: str, params: tuple = ()) -> sqlite3.Cursor:
        conn = self._get_connection()
        cursor = conn.cursor()
        cursor.execute(query, params)
        conn.commit()
        return cursor
    
    def close(self) -> None:
        if hasattr(self._local, "connection"):
            self._local.connection.close()
            del self._local.connection


def get_db() -> DatabaseConnection:
    return DatabaseConnection()
''',

    "app/database/repositories.py": '''
"""
Repository pattern for campaign data access.
"""

from typing import List, Optional, Dict, Any
from datetime import datetime, date

from .connection import DatabaseConnection
from ..models import Campaign, CampaignStatus, User


class CampaignRepository:
    """Repository for campaign CRUD operations."""
    
    def __init__(self, db: DatabaseConnection):
        self.db = db
    
    def _row_to_campaign(self, row: Dict[str, Any]) -> Campaign:
        return Campaign.from_dict(dict(row))
    
    def get_all(
        self, 
        user_id: int,
        status: Optional[str] = None,
        search: Optional[str] = None,
        sort_by: str = "start_date",
        sort_order: str = "ASC"
    ) -> List[Campaign]:
        query = "SELECT * FROM campaigns WHERE user_id = ?"
        params = [user_id]
        
        self._auto_complete_expired(user_id)
        
        if status:
            query += " AND status = ?"
            params.append(status)
        
        if search:
            query += """ AND (name LIKE ? OR description LIKE ? OR target_audience LIKE ? OR notes LIKE ? OR owner LIKE ? OR tags LIKE ?)"""
            search_pattern = f"%{search}%"
            params.extend([search_pattern] * 6)
        
        valid_sort_columns = {"id", "name", "budget", "spent", "start_date", "end_date", "status", "owner", "created_at", "updated_at"}
        if sort_by not in valid_sort_columns:
            sort_by = "start_date"
        
        sort_order = "ASC" if sort_order.upper() == "ASC" else "DESC"
        query += f" ORDER BY {sort_by} {sort_order}, id DESC"
        
        cursor = self.db.execute(query, tuple(params))
        rows = cursor.fetchall()
        return [self._row_to_campaign(dict(row)) for row in rows]
    
    def get_by_id(self, campaign_id: int, user_id: int) -> Optional[Campaign]:
        cursor = self.db.execute(
            "SELECT * FROM campaigns WHERE id = ? AND user_id = ?",
            (campaign_id, user_id)
        )
        row = cursor.fetchone()
        if not row:
            return None
        return self._row_to_campaign(dict(row))
    
    def create(self, campaign: Campaign, user_id: int) -> Campaign:
        now = datetime.now().isoformat()
        data = campaign.to_dict()
        data["user_id"] = user_id
        data["created_at"] = now
        data["updated_at"] = now
        
        cursor = self.db.execute("""
            INSERT INTO campaigns (
                user_id, name, description, budget, spent, currency, start_date,
                end_date, target_audience, status, owner, tags, assets,
                notes, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            user_id, data["name"], data["description"], data["budget"],
            data["spent"], data["currency"], data["start_date"],
            data["end_date"], data["target_audience"], data["status"],
            data["owner"], data["tags"], data["assets"],
            data["notes"], data["created_at"], data["updated_at"]
        ))
        
        campaign.id = cursor.lastrowid
        campaign.created_at = now
        campaign.updated_at = now
        return campaign
    
    def update(self, campaign_id: int, user_id: int, campaign: Campaign) -> Optional[Campaign]:
        existing = self.get_by_id(campaign_id, user_id)
        if not existing:
            return None
        
        now = datetime.now().isoformat()
        data = campaign.to_dict()
        data["updated_at"] = now
        
        self.db.execute("""
            UPDATE campaigns SET
                name = ?, description = ?, budget = ?, spent = ?,
                currency = ?, start_date = ?, end_date = ?,
                target_audience = ?, status = ?, owner = ?,
                tags = ?, assets = ?, notes = ?, updated_at = ?
            WHERE id = ? AND user_id = ?
        """, (
            data["name"], data["description"], data["budget"],
            data["spent"], data["currency"], data["start_date"],
            data["end_date"], data["target_audience"], data["status"],
            data["owner"], data["tags"], data["assets"],
            data["notes"], data["updated_at"], campaign_id, user_id
        ))
        
        campaign.id = campaign_id
        campaign.updated_at = now
        campaign.created_at = existing.created_at
        return campaign
    
    def delete(self, campaign_id: int, user_id: int) -> bool:
        cursor = self.db.execute(
            "DELETE FROM campaigns WHERE id = ? AND user_id = ?",
            (campaign_id, user_id)
        )
        return cursor.rowcount > 0
    
    def duplicate(self, campaign_id: int, user_id: int) -> Optional[Campaign]:
        existing = self.get_by_id(campaign_id, user_id)
        if not existing:
            return None
        
        new_campaign = Campaign(
            name=f"{existing.name} (Copy)",
            description=existing.description,
            budget=existing.budget,
            spent=0.0,
            currency=existing.currency,
            start_date=existing.start_date,
            end_date=existing.end_date,
            target_audience=existing.target_audience,
            status=CampaignStatus.DRAFT,
            owner=existing.owner,
            tags=existing.tags,
            assets=existing.assets,
            notes=existing.notes,
        )
        
        return self.create(new_campaign, user_id)
    
    def _auto_complete_expired(self, user_id: int) -> None:
        today = date.today().isoformat()
        self.db.execute(
            """
            UPDATE campaigns SET status = 'Completed', updated_at = ?
            WHERE user_id = ? AND (status = 'Active' OR status = 'Paused')
            AND end_date IS NOT NULL AND end_date < ?
            """,
            (datetime.now().isoformat(), user_id, today)
        )
    
    def get_stats(self, user_id: int) -> Dict[str, Any]:
        campaigns = self.get_all(user_id=user_id)
        
        stats = {
            "total": len(campaigns),
            "by_status": {s.value: 0 for s in CampaignStatus},
            "total_budget_usd": 0.0,
            "total_spent_usd": 0.0,
            "active_budget_usd": 0.0,
            "expired_count": 0,
        }
        
        from ..services.currency_service import CurrencyService
        currency_service = CurrencyService()
        
        for campaign in campaigns:
            stats["by_status"][campaign.status.value] = \
                stats["by_status"].get(campaign.status.value, 0) + 1
            
            budget_usd = currency_service.to_usd(campaign.budget, campaign.currency)
            spent_usd = currency_service.to_usd(campaign.spent, campaign.currency)
            
            stats["total_budget_usd"] += budget_usd
            stats["total_spent_usd"] += spent_usd
            
            if campaign.status == CampaignStatus.ACTIVE:
                stats["active_budget_usd"] += budget_usd
            
            if campaign.is_expired and campaign.status != CampaignStatus.COMPLETED:
                stats["expired_count"] += 1
        
        return stats


class UserRepository:
    """Repository for user operations."""
    
    def __init__(self, db: DatabaseConnection):
        self.db = db
    
    def create(self, user: User) -> User:
        now = datetime.now().isoformat()
        cursor = self.db.execute("""
            INSERT INTO users (email, username, password_hash, full_name, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?)
        """, (user.email, user.username, user.password_hash, user.full_name, now, now))
        
        user.id = cursor.lastrowid
        user.created_at = now
        user.updated_at = now
        return user
    
    def get_by_username(self, username: str) -> Optional[User]:
        cursor = self.db.execute(
            "SELECT * FROM users WHERE username = ?",
            (username,)
        )
        row = cursor.fetchone()
        if not row:
            return None
        return self._row_to_user(row)
    
    def get_by_email(self, email: str) -> Optional[User]:
        cursor = self.db.execute(
            "SELECT * FROM users WHERE email = ?",
            (email,)
        )
        row = cursor.fetchone()
        if not row:
            return None
        return self._row_to_user(row)
    
    def get_by_id(self, user_id: int) -> Optional[User]:
        cursor = self.db.execute(
            "SELECT * FROM users WHERE id = ?",
            (user_id,)
        )
        row = cursor.fetchone()
        if not row:
            return None
        return self._row_to_user(row)
    
    def _row_to_user(self, row) -> User:
        return User(
            id=row["id"],
            email=row["email"],
            username=row["username"],
            password_hash=row["password_hash"],
            full_name=row["full_name"],
            is_active=bool(row["is_active"]),
            created_at=row["created_at"],
            updated_at=row["updated_at"],
        )
''',

    "app/models/__init__.py": '''"""Models module for Stratis backend."""

from .campaign import Campaign
from .enums import CampaignStatus, Currency
from .user import User

__all__ = ["Campaign", "CampaignStatus", "Currency", "User"]
''',

    "app/models/enums.py": '''"""Enums for Stratis backend."""

from enum import Enum


class CampaignStatus(str, Enum):
    DRAFT = "Draft"
    ACTIVE = "Active"
    PAUSED = "Paused"
    COMPLETED = "Completed"
    
    @classmethod
    def _missing_(cls, value):
        if isinstance(value, str):
            value_lower = value.lower()
            for member in cls:
                if member.value.lower() == value_lower:
                    return member
        return None


class Currency(str, Enum):
    USD = "USD"
    EUR = "EUR"
    GBP = "GBP"
    
    @classmethod
    def _missing_(cls, value):
        if isinstance(value, str):
            value_upper = value.upper()
            for member in cls:
                if member.value == value_upper:
                    return member
        return None
''',

    "app/models/campaign.py": '''"""Campaign ORM model."""

from dataclasses import dataclass
from datetime import date, datetime
from typing import Optional, List, Dict, Any

from .enums import CampaignStatus, Currency


@dataclass
class Campaign:
    """Campaign domain model."""
    
    id: Optional[int] = None
    name: str = ""
    description: Optional[str] = None
    budget: float = 0.0
    spent: float = 0.0
    currency: Currency = Currency.USD
    start_date: str = ""
    end_date: Optional[str] = None
    target_audience: Optional[str] = None
    status: CampaignStatus = CampaignStatus.DRAFT
    owner: Optional[str] = None
    tags: Optional[List[str]] = None
    assets: Optional[str] = None
    notes: Optional[str] = None
    created_at: Optional[str] = None
    updated_at: Optional[str] = None
    
    def to_dict(self) -> Dict[str, Any]:
        return {
            "id": self.id,
            "name": self.name,
            "description": self.description,
            "budget": self.budget,
            "spent": self.spent,
            "currency": self.currency.value if self.currency else "USD",
            "start_date": self.start_date,
            "end_date": self.end_date,
            "target_audience": self.target_audience,
            "status": self.status.value if self.status else "Draft",
            "owner": self.owner,
            "tags": ",".join(self.tags) if self.tags else None,
            "assets": self.assets,
            "notes": self.notes,
            "created_at": self.created_at,
            "updated_at": self.updated_at,
        }
    
    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> "Campaign":
        tags = data.get("tags")
        if isinstance(tags, str) and tags:
            tags = [t.strip() for t in tags.split(",") if t.strip()]
        elif not tags:
            tags = None
            
        status = data.get("status")
        if isinstance(status, str):
            status = CampaignStatus(status)
        
        currency = data.get("currency")
        if isinstance(currency, str):
            currency = Currency(currency)
        
        return cls(
            id=data.get("id"),
            name=data.get("name", ""),
            description=data.get("description"),
            budget=data.get("budget", 0.0),
            spent=data.get("spent", 0.0),
            currency=currency or Currency.USD,
            start_date=data.get("start_date", ""),
            end_date=data.get("end_date"),
            target_audience=data.get("target_audience"),
            status=status or CampaignStatus.DRAFT,
            owner=data.get("owner"),
            tags=tags,
            assets=data.get("assets"),
            notes=data.get("notes"),
            created_at=data.get("created_at"),
            updated_at=data.get("updated_at"),
        )
    
    @property
    def is_expired(self) -> bool:
        if not self.end_date:
            return False
        try:
            end = datetime.strptime(self.end_date, "%Y-%m-%d").date()
            return end < date.today() and self.status != CampaignStatus.COMPLETED
        except ValueError:
            return False
    
    @property
    def remaining(self) -> float:
        return max(self.budget - self.spent, 0.0)
    
    @property
    def progress_pct(self) -> float:
        if self.budget <= 0:
            return 0.0
        return min((self.spent / self.budget) * 100, 100.0)
''',

    "app/models/user.py": '''"""User model."""

from dataclasses import dataclass
from typing import Optional


@dataclass
class User:
    """User domain model."""
    
    id: Optional[int] = None
    email: str = ""
    username: str = ""
    password_hash: str = ""
    full_name: Optional[str] = None
    is_active: bool = True
    created_at: Optional[str] = None
    updated_at: Optional[str] = None
    
    def to_dict(self) -> dict:
        return {
            "id": self.id,
            "email": self.email,
            "username": self.username,
            "full_name": self.full_name,
            "is_active": self.is_active,
            "created_at": self.created_at,
            "updated_at": self.updated_at,
        }
''',

    "app/schemas/__init__.py": '''"""Pydantic schemas for Stratis backend."""

from .campaign import (
    CampaignBase,
    CampaignCreate,
    CampaignUpdate,
    CampaignResponse,
    CampaignFilters,
)
from .ai import BriefRequest, BriefResponse, InsightsRequest, InsightsResponse
from .dashboard import DashboardStats
from .auth import UserRegister, UserLogin, Token, UserResponse, PasswordChange

__all__ = [
    "CampaignBase",
    "CampaignCreate",
    "CampaignUpdate",
    "CampaignResponse",
    "CampaignFilters",
    "BriefRequest",
    "BriefResponse",
    "InsightsRequest",
    "InsightsResponse",
    "DashboardStats",
    "UserRegister",
    "UserLogin",
    "Token",
    "UserResponse",
    "PasswordChange",
]
''',

    "app/schemas/campaign.py": '''"""Campaign Pydantic schemas for validation."""

from typing import Optional, List
from datetime import datetime
from pydantic import BaseModel, Field, validator

from ..models.enums import CampaignStatus, Currency


class CampaignBase(BaseModel):
    """Base campaign schema."""
    
    name: str = Field(..., min_length=1, max_length=100)
    description: Optional[str] = Field(None, max_length=1000)
    budget: float = Field(..., gt=0)
    spent: float = Field(0.0, ge=0)
    currency: Currency = Currency.USD
    start_date: str = Field(..., pattern=r"^\\d{4}-\\d{2}-\\d{2}$")
    end_date: Optional[str] = Field(None, pattern=r"^\\d{4}-\\d{2}-\\d{2}$")
    target_audience: Optional[str] = Field(None, max_length=500)
    status: CampaignStatus = CampaignStatus.DRAFT
    owner: Optional[str] = Field(None, max_length=100)
    tags: Optional[List[str]] = Field(None, max_length=20)
    assets: Optional[str] = Field(None, max_length=1000)
    notes: Optional[str] = Field(None, max_length=1000)
    
    @validator("end_date")
    def validate_end_date(cls, v, values):
        if v and "start_date" in values:
            start = datetime.strptime(values["start_date"], "%Y-%m-%d").date()
            end = datetime.strptime(v, "%Y-%m-%d").date()
            if end < start:
                raise ValueError("End date must be after start date")
        return v
    
    @validator("budget")
    def validate_budget(cls, v):
        if v > 1_000_000_000:
            raise ValueError("Budget too large (max 1 billion)")
        return v
    
    @validator("spent")
    def validate_spent(cls, v, values):
        if "budget" in values and v > values["budget"]:
            raise ValueError("Spent cannot exceed budget")
        return v


class CampaignCreate(CampaignBase):
    pass


class CampaignUpdate(BaseModel):
    name: Optional[str] = Field(None, min_length=1, max_length=100)
    description: Optional[str] = Field(None, max_length=1000)
    budget: Optional[float] = Field(None, gt=0)
    spent: Optional[float] = Field(None, ge=0)
    currency: Optional[Currency] = None
    start_date: Optional[str] = Field(None, pattern=r"^\\d{4}-\\d{2}-\\d{2}$")
    end_date: Optional[str] = Field(None, pattern=r"^\\d{4}-\\d{2}-\\d{2}$")
    target_audience: Optional[str] = Field(None, max_length=500)
    status: Optional[CampaignStatus] = None
    owner: Optional[str] = Field(None, max_length=100)
    tags: Optional[List[str]] = Field(None, max_length=20)
    assets: Optional[str] = Field(None, max_length=1000)
    notes: Optional[str] = Field(None, max_length=1000)
    
    @validator("end_date")
    def validate_end_date(cls, v, values):
        if v and "start_date" in values and values["start_date"]:
            start = datetime.strptime(values["start_date"], "%Y-%m-%d").date()
            end = datetime.strptime(v, "%Y-%m-%d").date()
            if end < start:
                raise ValueError("End date must be after start date")
        return v


class CampaignResponse(CampaignBase):
    id: int
    budget_usd: float
    is_expired: bool
    remaining: float
    progress_pct: float
    created_at: Optional[str] = None
    updated_at: Optional[str] = None
    
    class Config:
        from_attributes = True


class CampaignFilters(BaseModel):
    status: Optional[CampaignStatus] = None
    search: Optional[str] = Field(None, max_length=100)
    sort_by: str = "start_date"
    sort_order: str = "ASC"
    owner: Optional[str] = None
    start_date_from: Optional[str] = Field(None, pattern=r"^\\d{4}-\\d{2}-\\d{2}$")
    start_date_to: Optional[str] = Field(None, pattern=r"^\\d{4}-\\d{2}-\\d{2}$")
''',

    "app/schemas/ai.py": '''"""AI-related Pydantic schemas."""

from typing import Optional, List, Dict, Any
from pydantic import BaseModel, Field


class BriefRequest(BaseModel):
    name: str = Field(..., min_length=1, max_length=100)
    target_audience: Optional[str] = Field(None, max_length=500)
    budget: Optional[float] = Field(None, gt=0)
    currency: Optional[str] = Field("USD", pattern=r"^(USD|EUR|GBP)$")
    
    class Config:
        json_schema_extra = {
            "example": {
                "name": "Q4 Product Launch",
                "target_audience": "Tech-savvy professionals aged 25-40",
                "budget": 50000,
                "currency": "USD"
            }
        }


class BriefResponse(BaseModel):
    description: str
    tags: List[str]
    notes: str
    
    class Config:
        json_schema_extra = {
            "example": {
                "description": "Q4 Product Launch campaign targeting tech-savvy professionals.",
                "tags": ["product-launch", "q4-2024", "professional"],
                "notes": "Key metrics: 25% increase in brand awareness, 15% conversion rate."
            }
        }


class InsightsRequest(BaseModel):
    campaigns: List[Dict[str, Any]] = Field(..., min_length=1)
    
    class Config:
        json_schema_extra = {
            "example": {
                "campaigns": [
                    {
                        "name": "Q4 Product Launch",
                        "status": "Active",
                        "budget": 50000,
                        "spent": 25000,
                        "progress_pct": 50.0,
                        "is_expired": False
                    }
                ]
            }
        }


class InsightsResponse(BaseModel):
    insights: str
    
    class Config:
        json_schema_extra = {
            "example": {
                "insights": "Portfolio is healthy with 3 active campaigns. One campaign is approaching 90% spend. Recommend reallocating budget from underperforming channels."
            }
        }
''',

    "app/schemas/dashboard.py": '''"""Dashboard Pydantic schemas."""

from typing import Dict
from pydantic import BaseModel, Field


class DashboardStats(BaseModel):
    active_budget_usd: float = Field(..., description="Total active budget in USD")
    total_budget_usd: float = Field(..., description="Total budget across all campaigns in USD")
    total_spent_usd: float = Field(..., description="Total spent across all campaigns in USD")
    counts_by_status: Dict[str, int] = Field(..., description="Campaign counts by status")
    expired_count: int = Field(..., description="Number of expired campaigns")
    total_campaigns: int = Field(..., description="Total number of campaigns")
    
    class Config:
        json_schema_extra = {
            "example": {
                "active_budget_usd": 125000.00,
                "total_budget_usd": 250000.00,
                "total_spent_usd": 75000.00,
                "counts_by_status": {
                    "Draft": 2,
                    "Active": 3,
                    "Paused": 1,
                    "Completed": 4
                },
                "expired_count": 1,
                "total_campaigns": 10
            }
        }
''',

    "app/schemas/auth.py": '''"""Authentication schemas."""

from typing import Optional
from pydantic import BaseModel, Field, validator


class UserRegister(BaseModel):
    email: str = Field(..., max_length=100)
    username: str = Field(..., min_length=3, max_length=50)
    password: str = Field(..., min_length=8)
    full_name: Optional[str] = Field(None, max_length=100)
    
    @validator("email")
    def validate_email(cls, v):
        if "@" not in v or "." not in v:
            raise ValueError("Invalid email format")
        return v.lower()


class UserLogin(BaseModel):
    username: str
    password: str


class Token(BaseModel):
    access_token: str
    token_type: str = "bearer"


class UserResponse(BaseModel):
    id: int
    email: str
    username: str
    full_name: Optional[str] = None
    is_active: bool
    
    class Config:
        from_attributes = True


class PasswordChange(BaseModel):
    old_password: str
    new_password: str = Field(..., min_length=8)
''',

    "app/services/__init__.py": '''"""Services module for business logic."""

from .campaign_service import CampaignService
from .ai_service import AIService
from .currency_service import CurrencyService
from .dashboard_service import DashboardService
from .auth_service import AuthService
from .analytics_service import AnalyticsService
from .export_service import ExportService

__all__ = [
    "CampaignService",
    "AIService",
    "CurrencyService",
    "DashboardService",
    "AuthService",
    "AnalyticsService",
    "ExportService",
]
''',

    "app/services/campaign_service.py": '''
"""
Campaign business logic service.
"""

from typing import List, Optional
from ..models import Campaign
from ..database.repositories import CampaignRepository
from ..schemas import CampaignCreate, CampaignUpdate, CampaignFilters
from ..exceptions import NotFoundError, ValidationError


class CampaignService:
    """Service for campaign business logic."""
    
    def __init__(self, repository: CampaignRepository):
        self.repository = repository
    
    def get_all(self, user_id: int, filters: Optional[CampaignFilters] = None) -> List[Campaign]:
        if filters:
            return self.repository.get_all(
                user_id=user_id,
                status=filters.status.value if filters.status else None,
                search=filters.search,
                sort_by=filters.sort_by,
                sort_order=filters.sort_order,
            )
        return self.repository.get_all(user_id=user_id)
    
    def get_by_id(self, campaign_id: int, user_id: int) -> Optional[Campaign]:
        return self.repository.get_by_id(campaign_id, user_id)
    
    def create(self, user_id: int, campaign_data: CampaignCreate) -> Campaign:
        if campaign_data.start_date and campaign_data.end_date:
            from datetime import datetime
            start = datetime.strptime(campaign_data.start_date, "%Y-%m-%d").date()
            end = datetime.strptime(campaign_data.end_date, "%Y-%m-%d").date()
            if end < start:
                raise ValidationError("End date must be after start date")
        
        if campaign_data.spent > campaign_data.budget:
            raise ValidationError("Spent cannot exceed budget")
        
        campaign = Campaign(
            name=campaign_data.name,
            description=campaign_data.description,
            budget=campaign_data.budget,
            spent=campaign_data.spent,
            currency=campaign_data.currency,
            start_date=campaign_data.start_date,
            end_date=campaign_data.end_date,
            target_audience=campaign_data.target_audience,
            status=campaign_data.status,
            owner=campaign_data.owner,
            tags=campaign_data.tags,
            assets=campaign_data.assets,
            notes=campaign_data.notes,
        )
        
        return self.repository.create(campaign, user_id)
    
    def update(self, campaign_id: int, user_id: int, campaign_data: CampaignUpdate) -> Optional[Campaign]:
        existing = self.repository.get_by_id(campaign_id, user_id)
        if not existing:
            return None
        
        update_dict = {
            "name": campaign_data.name or existing.name,
            "description": campaign_data.description if campaign_data.description is not None else existing.description,
            "budget": campaign_data.budget or existing.budget,
            "spent": campaign_data.spent if campaign_data.spent is not None else existing.spent,
            "currency": campaign_data.currency or existing.currency,
            "start_date": campaign_data.start_date or existing.start_date,
            "end_date": campaign_data.end_date if campaign_data.end_date is not None else existing.end_date,
            "target_audience": campaign_data.target_audience if campaign_data.target_audience is not None else existing.target_audience,
            "status": campaign_data.status or existing.status,
            "owner": campaign_data.owner if campaign_data.owner is not None else existing.owner,
            "tags": campaign_data.tags if campaign_data.tags is not None else existing.tags,
            "assets": campaign_data.assets if campaign_data.assets is not None else existing.assets,
            "notes": campaign_data.notes if campaign_data.notes is not None else existing.notes,
        }
        
        if update_dict["start_date"] and update_dict["end_date"]:
            from datetime import datetime
            start = datetime.strptime(update_dict["start_date"], "%Y-%m-%d").date()
            end = datetime.strptime(update_dict["end_date"], "%Y-%m-%d").date()
            if end < start:
                raise ValidationError("End date must be after start date")
        
        if update_dict["spent"] > update_dict["budget"]:
            raise ValidationError("Spent cannot exceed budget")
        
        updated_campaign = Campaign(**update_dict)
        return self.repository.update(campaign_id, user_id, updated_campaign)
    
    def delete(self, campaign_id: int, user_id: int) -> bool:
        existing = self.repository.get_by_id(campaign_id, user_id)
        if not existing:
            return False
        return self.repository.delete(campaign_id, user_id)
    
    def duplicate(self, campaign_id: int, user_id: int) -> Optional[Campaign]:
        return self.repository.duplicate(campaign_id, user_id)
    
    def to_response(self, campaign: Campaign) -> dict:
        from .currency_service import CurrencyService
        
        currency_service = CurrencyService()
        
        return {
            "id": campaign.id,
            "name": campaign.name,
            "description": campaign.description,
            "budget": campaign.budget,
            "spent": campaign.spent,
            "currency": campaign.currency.value,
            "budget_usd": currency_service.to_usd(campaign.budget, campaign.currency),
            "start_date": campaign.start_date,
            "end_date": campaign.end_date,
            "target_audience": campaign.target_audience,
            "status": campaign.status.value,
            "owner": campaign.owner,
            "tags": campaign.tags,
            "assets": campaign.assets,
            "notes": campaign.notes,
            "is_expired": campaign.is_expired,
            "remaining": campaign.remaining,
            "progress_pct": campaign.progress_pct,
            "created_at": campaign.created_at,
            "updated_at": campaign.updated_at,
        }
''',

    "app/services/currency_service.py": '''"""Currency conversion service."""

from typing import Dict
from ..models.enums import Currency


class CurrencyService:
    """Service for currency conversion and formatting."""
    
    CONVERSION_RATES: Dict[Currency, float] = {
        Currency.USD: 1.0,
        Currency.EUR: 1.08,
        Currency.GBP: 1.27,
    }
    
    CURRENCY_SYMBOLS: Dict[Currency, str] = {
        Currency.USD: "$",
        Currency.EUR: "€",
        Currency.GBP: "£",
    }
    
    @classmethod
    def to_usd(cls, amount: float, currency: Currency) -> float:
        rate = cls.CONVERSION_RATES.get(currency, 1.0)
        return round(amount * rate, 2)
    
    @classmethod
    def get_symbol(cls, currency: Currency) -> str:
        return cls.CURRENCY_SYMBOLS.get(currency, currency.value)
''',

    "app/services/dashboard_service.py": '''
"""
Dashboard service for campaign statistics.
"""

from ..database.repositories import CampaignRepository
from ..schemas import DashboardStats


class DashboardService:
    """Service for dashboard statistics."""
    
    def __init__(self, repository: CampaignRepository):
        self.repository = repository
    
    def get_stats(self, user_id: int) -> DashboardStats:
        stats = self.repository.get_stats(user_id)
        
        return DashboardStats(
            active_budget_usd=round(stats["active_budget_usd"], 2),
            total_budget_usd=round(stats["total_budget_usd"], 2),
            total_spent_usd=round(stats["total_spent_usd"], 2),
            counts_by_status=stats["by_status"],
            expired_count=stats["expired_count"],
            total_campaigns=stats["total"],
        )
''',

    "app/services/auth_service.py": '''"""Authentication service."""

import bcrypt
from datetime import datetime, timedelta
from typing import Optional
from jose import JWTError, jwt

from ..config import config
from ..database.repositories import UserRepository
from ..exceptions import ValidationError
from ..models import User


class AuthService:
    """Service for authentication."""
    
    def __init__(self, user_repository: UserRepository):
        self.user_repository = user_repository
    
    def hash_password(self, password: str) -> str:
        salt = bcrypt.gensalt()
        return bcrypt.hashpw(password.encode("utf-8"), salt).decode("utf-8")
    
    def verify_password(self, password: str, password_hash: str) -> bool:
        return bcrypt.checkpw(password.encode("utf-8"), password_hash.encode("utf-8"))
    
    def create_access_token(self, user_id: int, username: str) -> str:
        expire = datetime.utcnow() + timedelta(minutes=config.auth.access_token_expire_minutes)
        payload = {
            "sub": str(user_id),
            "username": username,
            "exp": expire,
        }
        return jwt.encode(payload, config.auth.secret_key, algorithm=config.auth.algorithm)
    
    def decode_token(self, token: str) -> Optional[dict]:
        try:
            payload = jwt.decode(token, config.auth.secret_key, algorithms=[config.auth.algorithm])
            return payload
        except JWTError:
            return None
    
    def register_user(self, email: str, username: str, password: str, full_name: Optional[str] = None) -> User:
        if self.user_repository.get_by_username(username):
            raise ValidationError("Username already taken")
        
        if self.user_repository.get_by_email(email):
            raise ValidationError("Email already registered")
        
        password_hash = self.hash_password(password)
        
        user = User(
            email=email,
            username=username,
            password_hash=password_hash,
            full_name=full_name,
        )
        
        return self.user_repository.create(user)
    
    def authenticate_user(self, username: str, password: str) -> Optional[User]:
        user = self.user_repository.get_by_username(username)
        if not user:
            return None
        
        if not self.verify_password(password, user.password_hash):
            return None
        
        return user
    
    def get_current_user(self, token: str) -> Optional[User]:
        payload = self.decode_token(token)
        if not payload:
            return None
        
        user_id = payload.get("sub")
        if not user_id:
            return None
        
        return self.user_repository.get_by_id(int(user_id))
''',

    "app/services/ai_service.py": '''
"""
AI service for integration with LLM providers.
"""

import json
import logging
from typing import Dict, Any, List, Optional
from urllib.error import URLError
from urllib.request import Request, urlopen

from ..config import config
from ..exceptions import AIError

logger = logging.getLogger(__name__)


class AIService:
    """Service for AI-powered features."""
    
    def __init__(self):
        self.url = config.ai.url
        self.model = config.ai.model
        self.timeout = config.ai.timeout
        self.enabled = config.ai.enabled
    
    def _call_ollama(self, prompt: str) -> str:
        if not self.enabled:
            raise AIError("AI features are disabled")
        
        payload = json.dumps({
            "model": self.model,
            "prompt": prompt,
            "stream": False
        }).encode("utf-8")
        
        try:
            req = Request(
                self.url,
                data=payload,
                headers={"Content-Type": "application/json"},
                method="POST"
            )
            
            with urlopen(req, timeout=self.timeout) as response:
                data = json.loads(response.read().decode("utf-8"))
                return data.get("response", "").strip()
                
        except URLError as e:
            logger.error(f"Ollama connection error: {e}")
            raise AIError(f"Cannot connect to Ollama: {e.reason}")
        except json.JSONDecodeError as e:
            logger.error(f"Invalid JSON response from Ollama: {e}")
            raise AIError("Invalid response from AI service")
        except Exception as e:
            logger.error(f"Unexpected error calling Ollama: {e}")
            raise AIError(f"AI service error: {str(e)}")
    
    def _parse_json_response(self, raw: str, default: Dict[str, Any]) -> Dict[str, Any]:
        try:
            clean = raw.strip()
            if clean.startswith("```"):
                lines = clean.split("\\n")
                if lines[0].startswith("```"):
                    lines = lines[1:]
                if lines and lines[-1].startswith("```"):
                    lines = lines[:-1]
                clean = "\\n".join(lines)
            return json.loads(clean.strip())
        except json.JSONDecodeError:
            return default
    
    def generate_brief(
        self,
        name: str,
        target_audience: Optional[str] = None,
        budget: Optional[float] = None,
        currency: str = "USD"
    ) -> Dict[str, Any]:
        budget_info = f" with a budget of {budget} {currency}" if budget else ""
        audience_info = f" targeting {target_audience}" if target_audience else ""
        
        prompt = f"""You are a marketing strategist. Generate a brief for a campaign called "{name}"{audience_info}{budget_info}.

Respond ONLY with a JSON object in this exact format, no extra text:
{{
  "description": "2-3 sentence campaign description",
  "tags": ["tag1", "tag2", "tag3"],
  "notes": "1-2 sentences on key goals or success metrics"
}}"""
        
        raw_response = self._call_ollama(prompt)
        result = self._parse_json_response(raw_response, {
            "description": raw_response,
            "tags": [],
            "notes": ""
        })
        
        return result
    
    def generate_insights(self, campaigns: List[Dict[str, Any]]) -> str:
        if not campaigns:
            return "No campaigns to analyze."
        
        summary_lines = []
        for c in campaigns[:20]:
            line = (
                f"- \"{c.get('name', 'Unnamed')}\" | Status: {c.get('status', 'Unknown')} | "
                f"Budget: {c.get('budget', 0)} {c.get('currency', 'USD')} | "
                f"Spent: {c.get('spent', 0)} | "
                f"Progress: {c.get('progress_pct', 0)}% | "
                f"Expired: {c.get('is_expired', False)}"
            )
            summary_lines.append(line)
        
        campaigns_text = "\\n".join(summary_lines)
        
        prompt = f"""You are a marketing analyst. Here is a summary of marketing campaigns:

{campaigns_text}

Write a short, direct health report (3-5 sentences). Mention:
- Any campaigns that are over budget or nearly spent
- Any expired campaigns still marked active
- Overall portfolio health
- One concrete recommendation

Be direct and specific. No bullet points, just plain sentences."""
        
        raw_response = self._call_ollama(prompt)
        return raw_response
''',

    "app/services/analytics_service.py": '''"""Analytics service."""

from typing import Dict, Any
from ..database.repositories import CampaignRepository
from ..services.currency_service import CurrencyService


class AnalyticsService:
    """Service for campaign analytics."""
    
    def __init__(self, repository: CampaignRepository):
        self.repository = repository
        self.currency_service = CurrencyService()
    
    def get_analytics(self, user_id: int) -> Dict[str, Any]:
        campaigns = self.repository.get_all(user_id=user_id)
        
        total_budget = 0.0
        active_budget = 0.0
        completed_budget = 0.0
        total_spent = 0.0
        total_remaining = 0.0
        status_counts = {"Draft": 0, "Active": 0, "Paused": 0, "Completed": 0}
        
        for campaign in campaigns:
            budget_usd = self.currency_service.to_usd(campaign.budget, campaign.currency)
            spent_usd = self.currency_service.to_usd(campaign.spent, campaign.currency)
            
            total_budget += budget_usd
            total_spent += spent_usd
            total_remaining += campaign.remaining
            
            if campaign.status.value == "Active":
                active_budget += budget_usd
            elif campaign.status.value == "Completed":
                completed_budget += budget_usd
            
            status_counts[campaign.status.value] = status_counts.get(campaign.status.value, 0) + 1
        
        avg_budget = total_budget / len(campaigns) if campaigns else 0
        utilization = (total_spent / total_budget * 100) if total_budget > 0 else 0
        
        budget_by_status = {}
        for campaign in campaigns:
            status = campaign.status.value
            budget_usd = self.currency_service.to_usd(campaign.budget, campaign.currency)
            budget_by_status[status] = budget_by_status.get(status, 0) + budget_usd
        
        return {
            "total_budget_usd": round(total_budget, 2),
            "active_budget_usd": round(active_budget, 2),
            "completed_budget_usd": round(completed_budget, 2),
            "total_spent_usd": round(total_spent, 2),
            "total_remaining_usd": round(total_remaining, 2),
            "average_budget_usd": round(avg_budget, 2),
            "budget_utilization_pct": round(utilization, 1),
            "campaign_count": len(campaigns),
            "status_counts": status_counts,
            "budget_by_status": {k: round(v, 2) for k, v in budget_by_status.items()},
        }
''',

    "app/services/export_service.py": '''"""Export service for campaign data."""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExportService:
    """Service for exporting campaigns."""
    
    @staticmethod
    def to_excel(campaigns):
        """Export campaigns to Excel."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Campaigns"
        
        headers = [
            "ID", "Name", "Description", "Budget", "Spent", "Currency",
            "Start Date", "End Date", "Target Audience", "Status",
            "Owner", "Tags", "Assets", "Notes", "Created At", "Updated At"
        ]
        
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="0ea5e9", end_color="0ea5e9", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
            ws.column_dimensions[get_column_letter(col)].width = 20
        
        for row_idx, campaign in enumerate(campaigns, 2):
            data = [
                campaign.id,
                campaign.name,
                campaign.description or "",
                campaign.budget,
                campaign.spent,
                campaign.currency.value,
                campaign.start_date,
                campaign.end_date or "",
                campaign.target_audience or "",
                campaign.status.value,
                campaign.owner or "",
                ", ".join(campaign.tags) if campaign.tags else "",
                campaign.assets or "",
                campaign.notes or "",
                campaign.created_at or "",
                campaign.updated_at or "",
            ]
            
            for col_idx, value in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border
                cell.alignment = Alignment(vertical="center")
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
''',

    "app/routers/__init__.py": '''"""API routers for Stratis backend."""

from . import export, auth, campaigns, ai, dashboard, analytics

__all__ = ["export", "auth", "campaigns", "ai", "dashboard", "analytics"]
''',

    "app/routers/auth.py": '''"""Authentication routes."""

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.security import OAuth2PasswordBearer, OAuth2PasswordRequestForm

from ..database import get_db, DatabaseConnection
from ..database.repositories import UserRepository
from ..services.auth_service import AuthService
from ..schemas.auth import UserRegister, Token, UserResponse
from ..exceptions import ValidationError

router = APIRouter(prefix="/auth", tags=["Authentication"])

oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/auth/login")


def get_auth_service(db: DatabaseConnection = Depends(get_db)) -> AuthService:
    user_repository = UserRepository(db)
    return AuthService(user_repository)


@router.post("/register", response_model=UserResponse, status_code=status.HTTP_201_CREATED)
async def register(
    user_data: UserRegister,
    service: AuthService = Depends(get_auth_service),
):
    try:
        user = service.register_user(
            email=user_data.email,
            username=user_data.username,
            password=user_data.password,
            full_name=user_data.full_name,
        )
        return UserResponse(**user.to_dict())
    except ValidationError as e:
        raise HTTPException(status_code=400, detail=e.message)


@router.post("/login", response_model=Token)
async def login(
    form_data: OAuth2PasswordRequestForm = Depends(),
    service: AuthService = Depends(get_auth_service),
):
    user = service.authenticate_user(form_data.username, form_data.password)
    if not user:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid username or password",
            headers={"WWW-Authenticate": "Bearer"},
        )
    
    access_token = service.create_access_token(user.id, user.username)
    return {"access_token": access_token, "token_type": "bearer"}


@router.get("/me", response_model=UserResponse)
async def get_current_user(
    token: str = Depends(oauth2_scheme),
    service: AuthService = Depends(get_auth_service),
):
    user = service.get_current_user(token)
    if not user:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid authentication credentials",
            headers={"WWW-Authenticate": "Bearer"},
        )
    return UserResponse(**user.to_dict())


@router.post("/logout")
async def logout():
    return {"message": "Logged out successfully"}
''',

    "app/routers/campaigns.py": '''
"""
Campaign CRUD routes.
"""

from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException, status

from ..database import get_db, DatabaseConnection, CampaignRepository
from ..services import CampaignService
from ..schemas import CampaignCreate, CampaignUpdate, CampaignResponse, CampaignFilters
from ..exceptions import NotFoundError, ValidationError
from ..middleware.auth import require_auth
from ..models import User

router = APIRouter(prefix="/campaigns", tags=["Campaigns"])


def get_campaign_service(
    db: DatabaseConnection = Depends(get_db)
) -> CampaignService:
    repository = CampaignRepository(db)
    return CampaignService(repository)


@router.get("", response_model=List[CampaignResponse])
async def get_campaigns(
    current_user: User = Depends(require_auth()),
    status: Optional[str] = None,
    search: Optional[str] = None,
    sort_by: str = "start_date",
    sort_order: str = "ASC",
    service: CampaignService = Depends(get_campaign_service),
):
    filters = CampaignFilters(
        status=status,
        search=search,
        sort_by=sort_by,
        sort_order=sort_order,
    )
    
    try:
        campaigns = service.get_all(current_user.id, filters)
        return [service.to_response(c) for c in campaigns]
    except ValidationError as e:
        raise HTTPException(status_code=400, detail=e.message)


@router.get("/{campaign_id}", response_model=CampaignResponse)
async def get_campaign(
    campaign_id: int,
    current_user: User = Depends(require_auth()),
    service: CampaignService = Depends(get_campaign_service),
):
    try:
        campaign = service.get_by_id(campaign_id, current_user.id)
        if not campaign:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Campaign not found"
            )
        return service.to_response(campaign)
    except NotFoundError as e:
        raise HTTPException(status_code=404, detail=e.message)


@router.post("", response_model=CampaignResponse, status_code=status.HTTP_201_CREATED)
async def create_campaign(
    campaign: CampaignCreate,
    current_user: User = Depends(require_auth()),
    service: CampaignService = Depends(get_campaign_service),
):
    try:
        created = service.create(current_user.id, campaign)
        return service.to_response(created)
    except ValidationError as e:
        raise HTTPException(status_code=400, detail=e.message)


@router.put("/{campaign_id}", response_model=CampaignResponse)
async def update_campaign(
    campaign_id: int,
    campaign: CampaignUpdate,
    current_user: User = Depends(require_auth()),
    service: CampaignService = Depends(get_campaign_service),
):
    try:
        updated = service.update(campaign_id, current_user.id, campaign)
        if not updated:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Campaign not found"
            )
        return service.to_response(updated)
    except ValidationError as e:
        raise HTTPException(status_code=400, detail=e.message)
    except NotFoundError as e:
        raise HTTPException(status_code=404, detail=e.message)


@router.delete("/{campaign_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_campaign(
    campaign_id: int,
    current_user: User = Depends(require_auth()),
    service: CampaignService = Depends(get_campaign_service),
):
    try:
        deleted = service.delete(campaign_id, current_user.id)
        if not deleted:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Campaign not found"
            )
    except NotFoundError as e:
        raise HTTPException(status_code=404, detail=e.message)


@router.post("/{campaign_id}/duplicate", response_model=CampaignResponse, status_code=status.HTTP_201_CREATED)
async def duplicate_campaign(
    campaign_id: int,
    current_user: User = Depends(require_auth()),
    service: CampaignService = Depends(get_campaign_service),
):
    try:
        duplicated = service.duplicate(campaign_id, current_user.id)
        if not duplicated:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Campaign not found"
            )
        return service.to_response(duplicated)
    except NotFoundError as e:
        raise HTTPException(status_code=404, detail=e.message)
''',

    "app/routers/ai.py": '''"""AI-powered routes."""

from fastapi import APIRouter, Depends, HTTPException, status

from ..services import AIService
from ..schemas import BriefRequest, BriefResponse, InsightsRequest, InsightsResponse
from ..exceptions import AIError
from ..middleware.auth import require_auth
from ..models import User

router = APIRouter(prefix="/ai", tags=["AI"])


def get_ai_service() -> AIService:
    return AIService()


@router.post("/brief", response_model=BriefResponse)
async def generate_brief(
    request: BriefRequest,
    current_user: User = Depends(require_auth()),
    service: AIService = Depends(get_ai_service),
):
    try:
        result = service.generate_brief(
            name=request.name,
            target_audience=request.target_audience,
            budget=request.budget,
            currency=request.currency,
        )
        return result
    except AIError as e:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail=f"AI service unavailable: {e.message}"
        )


@router.post("/insights", response_model=InsightsResponse)
async def generate_insights(
    request: InsightsRequest,
    current_user: User = Depends(require_auth()),
    service: AIService = Depends(get_ai_service),
):
    try:
        result = service.generate_insights(request.campaigns)
        return {"insights": result}
    except AIError as e:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail=f"AI service unavailable: {e.message}"
        )
''',

    "app/routers/dashboard.py": '''"""Dashboard routes."""

from fastapi import APIRouter, Depends

from ..database import get_db, DatabaseConnection, CampaignRepository
from ..services import DashboardService
from ..schemas import DashboardStats
from ..middleware.auth import require_auth
from ..models import User

router = APIRouter(prefix="/dashboard", tags=["Dashboard"])


def get_dashboard_service(
    db: DatabaseConnection = Depends(get_db)
) -> DashboardService:
    repository = CampaignRepository(db)
    return DashboardService(repository)


@router.get("/stats", response_model=DashboardStats)
async def get_dashboard_stats(
    current_user: User = Depends(require_auth()),
    service: DashboardService = Depends(get_dashboard_service),
):
    return service.get_stats(current_user.id)
''',

    "app/routers/analytics.py": '''"""Analytics routes."""

from fastapi import APIRouter, Depends

from ..database import get_db, DatabaseConnection, CampaignRepository
from ..services import AnalyticsService
from ..middleware.auth import require_auth
from ..models import User

router = APIRouter(prefix="/analytics", tags=["Analytics"])


def get_analytics_service(db: DatabaseConnection = Depends(get_db)) -> AnalyticsService:
    repository = CampaignRepository(db)
    return AnalyticsService(repository)


@router.get("/stats")
async def get_analytics(
    current_user: User = Depends(require_auth()),
    service: AnalyticsService = Depends(get_analytics_service),
):
    return service.get_analytics(current_user.id)
''',

    "app/routers/export.py": '''"""Export routes for campaign data."""

import csv
import io
from datetime import date
from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse

from ..database import get_db, DatabaseConnection, CampaignRepository
from ..services import CampaignService, ExportService
from ..middleware.auth import require_auth
from ..models import User

router = APIRouter(prefix="/export", tags=["Export"])


def get_campaign_service(
    db: DatabaseConnection = Depends(get_db)
) -> CampaignService:
    repository = CampaignRepository(db)
    return CampaignService(repository)


@router.get("/csv", response_class=StreamingResponse)
async def export_campaigns_csv(
    current_user: User = Depends(require_auth()),
    service: CampaignService = Depends(get_campaign_service),
):
    campaigns = service.get_all(current_user.id)
    
    output = io.StringIO()
    fieldnames = [
        "id", "name", "description", "budget", "spent", "currency",
        "start_date", "end_date", "target_audience", "status",
        "owner", "tags", "assets", "notes", "created_at", "updated_at"
    ]
    
    writer = csv.DictWriter(output, fieldnames=fieldnames)
    writer.writeheader()
    
    for campaign in campaigns:
        row = campaign.to_dict()
        if row.get("tags"):
            row["tags"] = ", ".join(campaign.tags) if campaign.tags else ""
        writer.writerow(row)
    
    output.seek(0)
    filename = f"stratis_campaigns_{date.today().isoformat()}.csv"
    
    async def iter_csv():
        yield output.getvalue()
    
    return StreamingResponse(
        iter_csv(),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )


@router.get("/excel", response_class=StreamingResponse)
async def export_campaigns_excel(
    current_user: User = Depends(require_auth()),
    service: CampaignService = Depends(get_campaign_service),
):
    campaigns = service.get_all(current_user.id)
    
    excel_file = ExportService.to_excel(campaigns)
    filename = f"stratis_campaigns_{date.today().isoformat()}.xlsx"
    
    async def iter_excel():
        yield excel_file.getvalue()
    
    return StreamingResponse(
        iter_excel(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
''',

    "app/middleware/__init__.py": '''"""Middleware for Stratis backend."""

from .auth import require_auth, get_current_user_from_request

__all__ = ["require_auth", "get_current_user_from_request"]
''',

    "app/middleware/auth.py": '''"""Authentication middleware."""

from fastapi import Request, HTTPException, status, Depends
from fastapi.security import OAuth2PasswordBearer

from ..database import get_db, DatabaseConnection
from ..database.repositories import UserRepository
from ..services.auth_service import AuthService
from ..models import User

oauth2_scheme = OAuth2PasswordBearer(tokenUrl="/auth/login")


def get_auth_service(db: DatabaseConnection = Depends(get_db)) -> AuthService:
    user_repository = UserRepository(db)
    return AuthService(user_repository)


def require_auth():
    """Dependency for protected routes."""
    async def dependency(token: str = Depends(oauth2_scheme), db: DatabaseConnection = Depends(get_db)):
        user_repository = UserRepository(db)
        auth_service = AuthService(user_repository)
        user = auth_service.get_current_user(token)
        if not user:
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail="Invalid authentication credentials",
                headers={"WWW-Authenticate": "Bearer"},
            )
        return user
    return dependency


async def get_current_user_from_request(request: Request) -> User:
    auth_header = request.headers.get("Authorization")
    if not auth_header or not auth_header.startswith("Bearer "):
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Missing or invalid authorization header",
        )
    
    token = auth_header.split(" ")[1]
    db = get_db()
    user_repository = UserRepository(db)
    auth_service = AuthService(user_repository)
    user = auth_service.get_current_user(token)
    
    if not user:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid authentication credentials",
        )
    
    return user
''',

    "app/exceptions/__init__.py": '''"""Exception handling for Stratis backend."""

from .base import (
    StratisError,
    DatabaseError,
    NotFoundError,
    ValidationError,
    AIError,
    ConfigurationError,
)
from .handlers import (
    validation_exception_handler,
    http_exception_handler,
    generic_exception_handler,
    ai_exception_handler,
)

__all__ = [
    "StratisError",
    "DatabaseError",
    "NotFoundError",
    "ValidationError",
    "AIError",
    "ConfigurationError",
    "validation_exception_handler",
    "http_exception_handler",
    "generic_exception_handler",
    "ai_exception_handler",
]
''',

    "app/exceptions/base.py": '''"""Base exception classes for Stratis."""

from typing import Optional, Any


class StratisError(Exception):
    """Base exception for all Stratis errors."""
    
    def __init__(self, message: str, detail: Optional[Any] = None):
        self.message = message
        self.detail = detail
        super().__init__(message)


class DatabaseError(StratisError):
    pass


class NotFoundError(StratisError):
    pass


class ValidationError(StratisError):
    pass


class AIError(StratisError):
    pass


class ConfigurationError(StratisError):
    pass
''',

    "app/exceptions/handlers.py": '''"""Exception handlers for FastAPI."""

import logging
from fastapi import Request, HTTPException
from fastapi.responses import JSONResponse
from fastapi.exceptions import RequestValidationError

from .base import AIError

logger = logging.getLogger(__name__)


async def validation_exception_handler(
    request: Request, 
    exc: RequestValidationError
) -> JSONResponse:
    errors = []
    for error in exc.errors():
        field = " -> ".join(str(loc) for loc in error["loc"])
        errors.append({
            "field": field,
            "message": error["msg"],
            "type": error["type"],
        })
    
    logger.warning(f"Validation error: {errors}")
    
    return JSONResponse(
        status_code=422,
        content={
            "detail": "Validation error",
            "errors": errors,
        }
    )


async def http_exception_handler(
    request: Request,
    exc: HTTPException
) -> JSONResponse:
    logger.warning(f"HTTP {exc.status_code}: {exc.detail}")
    return JSONResponse(
        status_code=exc.status_code,
        content={"detail": exc.detail},
    )


async def generic_exception_handler(
    request: Request,
    exc: Exception
) -> JSONResponse:
    logger.error(f"Unhandled exception: {exc}", exc_info=True)
    return JSONResponse(
        status_code=500,
        content={
            "detail": "Internal server error",
        }
    )


async def ai_exception_handler(
    request: Request,
    exc: AIError
) -> JSONResponse:
    logger.warning(f"AI error: {exc.message}")
    return JSONResponse(
        status_code=503,
        content={
            "detail": exc.message,
            "hint": "Make sure Ollama is running with 'ollama serve'",
        }
    )
''',

    "app/utils/__init__.py": '''"""Utilities for Stratis backend."""

from .validators import (
    validate_email,
    validate_date_range,
    sanitize_string,
    truncate_string,
)

__all__ = [
    "validate_email",
    "validate_date_range",
    "sanitize_string",
    "truncate_string",
]
''',

    "app/utils/validators.py": '''"""Custom validation utilities."""

import re
from datetime import datetime
from typing import Optional, Tuple


def validate_email(email: str) -> bool:
    if not email:
        return True
    pattern = r"^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\\.[a-zA-Z]{2,}$"
    return bool(re.match(pattern, email))


def validate_date_range(
    start_date: str,
    end_date: Optional[str]
) -> Tuple[bool, Optional[str]]:
    if not end_date:
        return True, None
    
    try:
        start = datetime.strptime(start_date, "%Y-%m-%d").date()
        end = datetime.strptime(end_date, "%Y-%m-%d").date()
        if end < start:
            return False, "End date must be after start date"
        return True, None
    except ValueError:
        return False, "Invalid date format. Use YYYY-MM-DD"


def sanitize_string(text: str) -> str:
    if not text:
        return text
    return ''.join(char for char in text if ord(char) >= 32 or char == '\\n')


def truncate_string(text: str, max_length: int) -> str:
    if not text or len(text) <= max_length:
        return text
    return text[:max_length - 3] + "..."
''',

    "app/main.py": '''
"""
Stratis Marketing Engine - Main Application
"""

import logging
from contextlib import asynccontextmanager
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.exceptions import RequestValidationError

from .config import config
from .routers import auth, export, campaigns, ai, dashboard, analytics
from .exceptions.handlers import (
    validation_exception_handler,
    http_exception_handler,
    generic_exception_handler,
    ai_exception_handler,
)
from .exceptions import AIError

logging.basicConfig(
    level=logging.INFO if not config.api.debug else logging.DEBUG,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
)
logger = logging.getLogger(__name__)


@asynccontextmanager
async def lifespan(app: FastAPI):
    logger.info("Starting Stratis Marketing Engine...")
    config.ensure_directories()
    logger.info(f"Database: {config.database.absolute_path}")
    logger.info(f"Environment: {config.api.environment}")
    logger.info(f"AI Enabled: {config.ai.enabled}")
    yield
    logger.info("Shutting down Stratis Marketing Engine...")


def create_app() -> FastAPI:
    app = FastAPI(
        title="Stratis Marketing Engine",
        description="Marketing campaign management with AI capabilities",
        version="2.0.0",
        lifespan=lifespan,
        debug=config.api.debug,
        docs_url="/docs",
        redoc_url="/redoc",
    )
    
    app.add_middleware(
        CORSMiddleware,
        allow_origins=config.cors.origins,
        allow_credentials=config.cors.allow_credentials,
        allow_methods=config.cors.allow_methods,
        allow_headers=config.cors.allow_headers,
    )
    
    app.add_exception_handler(RequestValidationError, validation_exception_handler)
    app.add_exception_handler(HTTPException, http_exception_handler)
    app.add_exception_handler(Exception, generic_exception_handler)
    app.add_exception_handler(AIError, ai_exception_handler)
    
    api_prefix = config.api.prefix
    
    app.include_router(auth.router, prefix=api_prefix)
    app.include_router(export.router, prefix=api_prefix)  # /export/*
    app.include_router(campaigns.router, prefix=api_prefix)
    app.include_router(ai.router, prefix=api_prefix)
    app.include_router(dashboard.router, prefix=api_prefix)
    app.include_router(analytics.router, prefix=api_prefix)
    
    @app.get("/", tags=["Health"])
    async def root():
        return {
            "message": "Stratis API Engine Operating Nominally.",
            "version": "2.0.0",
            "environment": config.api.environment,
        }
    
    return app


app = create_app()
''',

    "run.py": '''#!/usr/bin/env python
"""
Development server runner.
"""

import uvicorn

from app.config import config

if __name__ == "__main__":
    uvicorn.run(
        "app.main:app",
        host=config.api.host,
        port=config.api.port,
        reload=config.api.reload,
        log_level="info",
    )
''',

    "requirements.txt": '''fastapi==0.104.1
uvicorn[standard]==0.24.0
pydantic==2.5.0
bcrypt==4.1.2
python-jose[cryptography]==3.3.0
python-multipart==0.0.6
passlib[bcrypt]==1.7.4
openpyxl==3.1.2
pytest==7.4.3
pytest-asyncio==0.21.1
httpx==0.25.1
''',

    ".env.example": '''# Database
STRATIS_DB_PATH=campaigns.db

# API
STRATIS_ENVIRONMENT=development
STRATIS_DEBUG=true
STRATIS_HOST=127.0.0.1
STRATIS_PORT=8000
STRATIS_RELOAD=true

# AI
STRATIS_OLLAMA_URL=http://localhost:11434/api/generate
STRATIS_OLLAMA_MODEL=llama3.2
STRATIS_OLLAMA_TIMEOUT=60
STRATIS_AI_ENABLED=true

# CORS
STRATIS_CORS_ORIGINS=http://localhost:5500,http://127.0.0.1:5500

# Rate Limiting
STRATIS_RATE_LIMIT_REQUESTS=100
STRATIS_RATE_LIMIT_PERIOD=60

# Authentication
STRATIS_SECRET_KEY=dev-secret-key-change-me
STRATIS_JWT_ALGORITHM=HS256
STRATIS_ACCESS_TOKEN_EXPIRE_MINUTES=30
''',
}

# Create directories and files
for filepath, content in files.items():
    # Create directory if needed
    dirname = os.path.dirname(filepath)
    if dirname:
        os.makedirs(dirname, exist_ok=True)
    
    # Write file
    with open(filepath, 'w', encoding='utf-8') as f:
        f.write(content.lstrip())
        print(f"Created: {filepath}")

print("\n✅ All files recreated successfully!")
print("\nNext steps:")
print("1. python -m pip install -r requirements.txt")
print("2. python run.py")